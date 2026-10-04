"""Export tables/folders to XLSX (with thumbnails) and CSV."""
import csv
import io
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

from .. import quantity
from ..extensions import db
from ..models import InvTable, User
from . import photos as photo_svc
from . import settings
from .tree import FolderMap

THUMB = 90
_BAD = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(name: str, used: set[str]) -> str:
    base = _BAD.sub(" ", name).strip()[:31] or "Sheet"
    out, i = base, 2
    while out.lower() in used:
        suffix = f" ({i})"
        out = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(out.lower())
    return out


def _cell_value(item, col):
    key = col["key"]
    if key == "description":
        return item.description
    if key == "quantity":
        return quantity.format(item.quantity)
    if key == "observation":
        return item.observation
    if key == "photos":
        return None
    v = (item.custom or {}).get(key)
    if col["type"] == "checkbox":
        return "Yes" if v else ""
    return v


def _visible_columns(t: InvTable):
    return [c for c in t.columns if not c.get("hidden")]


def _thumb_bytes(sha: str) -> io.BytesIO:
    img = Image.open(photo_svc.path_for(sha, "thumb"))
    img.thumbnail((THUMB, THUMB))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=80)
    buf.seek(0)
    return buf


def _write_table(ws, t: InvTable, fmap: FolderMap):
    accent = (settings.get("app.primary_color") or "#4f46e5").lstrip("#")
    updater = db.session.get(User, t.content_updated_by) if t.content_updated_by else None
    path = " / ".join(fmap.path(t.folder_id))
    cols = _visible_columns(t)
    header_lines = [
        settings.get("app.name"),
        f"Location: {path}" if path else None,
        f"Table: {t.name}",
        f"Summary: {t.summary}" if t.summary else None,
        "Last update: " + settings.local(t.content_updated_at).strftime("%Y-%m-%d %H:%M %Z")
        + (f" by {updater.name}" if updater else ""),
        "Archived" if not (t.active and fmap.effective_active(t.folder_id)) else None,
    ]
    r = 1
    for line in [x for x in header_lines if x]:
        ws.cell(row=r, column=1, value=line).font = Font(bold=(r == 1), size=13 if r == 1 else 11)
        if len(cols) > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        r += 1
    r += 1
    header_row = r
    for ci, col in enumerate(cols, start=1):
        c = ws.cell(row=r, column=ci, value=col.get("label") or col["key"])
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=accent)
        ws.column_dimensions[get_column_letter(ci)].width = max(10, min(70, (col.get("width") or 150) / 7))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    extra = len(cols) + 1  # "Added" column after the table's own columns
    c = ws.cell(row=header_row, column=extra, value="Added")
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=accent)
    ws.column_dimensions[get_column_letter(extra)].width = 24
    photo_ci = next((i for i, c in enumerate(cols, start=1) if c["key"] == "photos"), None)
    for item in t.live_items():
        r += 1
        for ci, col in enumerate(cols, start=1):
            c = ws.cell(row=r, column=ci, value=_cell_value(item, col))
            c.alignment = Alignment(wrap_text=True, vertical="top")
        added = settings.local(item.created_at).strftime("%Y-%m-%d %H:%M")
        if item.creator:
            added += f" by {item.creator.name}"
        ws.cell(row=r, column=extra, value=added).alignment = Alignment(vertical="top")
        if photo_ci and item.photos:
            try:
                xl = XLImage(_thumb_bytes(item.photos[0].sha256))
                ws.add_image(xl, f"{get_column_letter(photo_ci)}{r}")
                ws.row_dimensions[r].height = THUMB * 0.78
            except FileNotFoundError:
                pass
    if photo_ci:
        ws.column_dimensions[get_column_letter(photo_ci)].width = 15


def tables_for(folder_id: int | None = None, table_ids: list[int] | None = None,
               include_archived=True) -> list[InvTable]:
    fmap = FolderMap()
    q = db.session.query(InvTable).filter(InvTable.deleted_at.is_(None))
    if table_ids:
        q = q.filter(InvTable.id.in_(table_ids))
    elif folder_id is not None:
        q = q.filter(InvTable.folder_id.in_([folder_id] + fmap.descendants(folder_id)))
    out = []
    for t in q.order_by(InvTable.position, InvTable.name):
        if fmap.effective_deleted(t.folder_id):
            continue
        if not include_archived and not (t.active and fmap.effective_active(t.folder_id)):
            continue
        out.append(t)
    out.sort(key=lambda t: (fmap.path(t.folder_id), t.position, t.name))
    return out


def to_xlsx(tables: list[InvTable]) -> bytes:
    fmap = FolderMap()
    wb = Workbook()
    used: set[str] = set()
    if len(tables) > 1:
        idx = wb.active
        idx.title = _sheet_name("Index", used)
        idx.append([settings.get("app.name")])
        idx["A1"].font = Font(bold=True, size=13)
        idx.append([])
        idx.append(["Table", "Location", "Summary", "Items", "Last update"])
        for c in idx[3]:
            c.font = Font(bold=True)
        names = []
        for t in tables:
            names.append(_sheet_name(t.name, used))
        for t, sname in zip(tables, names):
            idx.append([f'=HYPERLINK("#\'{sname}\'!A1","{t.name.replace(chr(34), "")}")',
                        " / ".join(fmap.path(t.folder_id)), t.summary or "", t.live_item_count(),
                        settings.local(t.content_updated_at).strftime("%Y-%m-%d")])
        for col, w in zip("ABCDE", (40, 30, 50, 8, 14)):
            idx.column_dimensions[col].width = w
        for t, sname in zip(tables, names):
            _write_table(wb.create_sheet(sname), t, fmap)
    else:
        ws = wb.active
        ws.title = _sheet_name(tables[0].name, used)
        _write_table(ws, tables[0], fmap)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(t: InvTable) -> str:
    cols = [c for c in _visible_columns(t) if c["key"] != "photos"]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c.get("label") or c["key"] for c in cols] + ["Estimated", "Added at", "Added by"])
    for item in t.live_items():
        w.writerow([_cell_value(item, c) or "" for c in cols] + [
            "yes" if item.qty_estimated else "",
            settings.local(item.created_at).strftime("%Y-%m-%d %H:%M"),
            item.creator.name if item.creator else ""])
    return buf.getvalue()
