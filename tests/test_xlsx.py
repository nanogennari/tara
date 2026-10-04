import io

from openpyxl import load_workbook

from sample_workbook import build, stats


def _upload(client, path, **form):
    data = {"file": (io.BytesIO(build()), "Sample Inventory.xlsx"), **form}
    return client.post(path, data=data, content_type="multipart/form-data")


def test_analyze_detects_sheets(client):
    sheets = _upload(client, "/api/import/xlsx/analyze").get_json()
    by = {s["sheet"]: s for s in sheets}
    assert by["Index"]["include"] is False  # table of contents, not inventory
    med = by["Box 01 Medicine Box"]
    assert med["include"] and med["rows"] == 12 and med["images"] == 12
    assert med["name"] == "Box 01: Medicine Box"
    assert med["summary"] == "Assorted medicine and Band-Aids"
    maps = {c["header"]: (c["maps_to"], c.get("type")) for c in med["columns"]}
    assert maps["Item"] == ("description", None)
    assert maps["Expiry"] == ("custom", "date")
    assert maps["Bag"] == ("custom", "select")


def test_import_sample_inventory(client):
    f = client.post("/api/folders", json={"name": "Storage"}).get_json()
    res = _upload(client, "/api/import/xlsx", folder_id=str(f["id"])).get_json()
    expected = stats()
    assert len(res["tables"]) == expected["tables"]
    assert res["items"] == expected["items"]
    assert res["photos"] == expected["photos"]

    tree = client.get("/api/tree").get_json()
    med = next(t for t in tree["tables"] if t["name"].startswith("Box 01"))
    data = client.get(f"/api/tables/{med['id']}").get_json()
    assert "Location: Example Storage Room" in data["context"]
    cols = {c["key"]: c for c in data["columns"]}
    assert cols["expiry"]["type"] == "date"
    assert cols["bag"]["options"] == ["Teal bag", "Red bag", "Loose"]
    neo = data["items"][0]
    assert neo["description"] == "Neosporin ointment"
    assert neo["custom"] == {"bag": "Teal bag", "expiry": "2026-04-30"}
    assert len(neo["photos"]) == 1
    valerian = next(i for i in data["items"] if i["description"] == "Valerian extract")
    assert "expiry" not in valerian["custom"]  # "No date" -> empty

    zome = next(t for t in tree["tables"] if t["name"].startswith("Box 04"))
    first = client.get(f"/api/tables/{zome['id']}").get_json()["items"][0]
    assert first["quantity"] == {"type": "count", "value": 75.0, "unit": "", "text": "", "estimated": True}


def test_export_roundtrip(client):
    _upload(client, "/api/import/xlsx")
    r = client.get("/api/export/xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames[0] == "Index" and len(wb.sheetnames) == stats()["tables"] + 1
    ws = wb[wb.sheetnames[1]]
    assert any(str(c.value or "").startswith("Last update:") for c in ws["A"][:8])
    assert len(ws._images) == 12

    tid = client.get("/api/tree").get_json()["tables"][0]["id"]
    csv = client.get(f"/api/tables/{tid}/export.csv").get_data(as_text=True)
    assert "Neosporin ointment" in csv
