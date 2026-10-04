"""Synthetic inventory workbook for tests and demos (no real data).

Mirrors the layout of a typical hand-made inventory spreadsheet:
an index sheet of HYPERLINKs, title rows ("Location: ...", "Summary: ..."),
a header row, a choice-like column, an expiry date column with "No date"
values, estimated quantities ("~75") and one embedded photo per row.

    python tests/sample_workbook.py out.xlsx
"""
import io
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image, ImageDraw

ORG = "Example Camp Inventory"
LOCATION = "Location: Example Storage Room"

SHEETS = [
    {
        "sheet": "Box 01 Medicine Box", "name": "Box 01: Medicine Box", "summary": "Assorted medicine and Band-Aids",
        "header": ["Bag", "Item", "Qty", "Expiry", "Photo"],
        "rows": [
            ["Teal bag", "Neosporin ointment", "1 tube, 14.2 g (0.5 oz)", datetime(2026, 4, 30)],
            ["Teal bag", "Aleve (naproxen 220 mg)", "90 tabs", datetime(2026, 3, 31)],
            ["Teal bag", "Imodium (loperamide 2 mg)", "8 caplets", datetime(2026, 12, 31)],
            ["Teal bag", "Advil Liqui-Gels minis (ibuprofen 200 mg)", "80 caps", datetime(2028, 1, 31)],
            ["Teal bag", "Tylenol drops (paracetamol 200 mg/mL)", "15 mL", datetime(2028, 10, 31)],
            ["Teal bag", "Valerian extract", "1 fl oz", "No date"],
            ["Teal bag", "Band-Aid, large", "12", "No date"],
            ["Teal bag", "Band-Aid, medium", "7", "No date"],
            ["Red bag", "Bandages, small (some latex)", "30+", "No date"],
            ["Red bag", "Self-adhesive wrap, red", "1", "No date"],
            ["Red bag", "Excedrin Migraine", "?", datetime(2028, 2, 29)],
            ["Loose", "DayQuil + NyQuil Severe Cold & Flu", "16 + 8 caps", datetime(2028, 1, 31)],
        ],
    },
    {
        "sheet": "Box 02 Laptops", "name": "Box 02: Laptops and Chargers", "summary": "Laptops and chargers",
        "header": ["Item", "Qty", "Obs", "Photo"],
        "rows": [
            ["Lenovo ThinkPad X13, AMD Ryzen 5, 16 GB RAM", "8", "Labeled LAB-02 to LAB-09"],
            ["Laptop charger, USB-C 65 W, with AC power cord", "10", "Labeled LAB-01 to LAB-10"],
            ["Extension cord, 3 m, 3 outlets (sealed)", "3", "Cheap, not good quality"],
        ],
    },
    {
        "sheet": "Box 03 Office Supplies", "name": "Box 03: Office Supplies", "summary": "Markers, name tags, tape and paper",
        "header": ["Item", "Qty", "Obs", "Photo"],
        "rows": [
            ["Sharpie Fine Point permanent marker, black", "~22", "Approximate count, in ziplock"],
            ["Sharpie Fine Point permanent marker, assorted colors", "~24", "Approximate count, in ziplock"],
            ["Magnetic name tags, assorted colors", "~70", "Rough estimate, in clear box"],
            ["Scotch clear packing tape", "3 rolls", "1 with red dispenser"],
            ["Duck duct tape, silver, 48 mm x 55 m", "2 rolls", "1 sealed, 1 opened"],
            ["Scissors, 21 cm, stainless steel", "6", ""],
            ["Glue sticks, 40 g", "1 pack (12)", "Unopened"],
            ["Avery mini business cards, 1\" x 3\"", "3 packs", "160 cards each"],
            ["Origami paper, 15 cm, double-sided", "1 pack (500 sheets)", ""],
        ],
    },
    {
        "sheet": "Box 04 Zometool", "name": "Box 04: Zometool", "summary": "Zometool struts and connectors",
        "header": ["Item", "Qty", "Obs", "Photo"],
        "rows": [
            ["Wooden cat tessellation pieces (3 bags)", "~75", "Estimated by weight (108 g total)"],
            ["Zometool blue struts, bag A (2 ziplocks)", "~370", "Estimated by weight (871 g)"],
            ["Zometool white connector balls, cloth drawstring bag", "~1,850", "Estimated by weight (1.99 kg)"],
        ],
    },
]

COLORS = ["#e11d48", "#2563eb", "#16a34a", "#ca8a04", "#7c3aed", "#0891b2", "#ea580c", "#475569"]


def _photo(label: str, i: int) -> io.BytesIO:
    img = Image.new("RGB", (160, 120), COLORS[i % len(COLORS)])
    ImageDraw.Draw(img).text((8, 52), label[:22], fill="white")
    buf = io.BytesIO()
    img.save(buf, "JPEG")
    buf.seek(0)
    return buf


def stats() -> dict:
    rows = sum(len(s["rows"]) for s in SHEETS)
    return {"tables": len(SHEETS), "items": rows, "photos": rows}


def build(path=None) -> bytes:
    wb = Workbook()
    idx = wb.active
    idx.title = "Index"
    idx.append([ORG])
    idx.append([LOCATION])
    idx.append([])
    idx.append(["Container", "Summary"])
    for s in SHEETS:
        idx.append([f'=HYPERLINK("#\'{s["sheet"]}\'!A1","{s["name"]}")', s["summary"]])

    n = 0
    for s in SHEETS:
        ws = wb.create_sheet(s["sheet"])
        ws.append([ORG])
        ws.append([LOCATION])
        ws.append([f"Box: {s['name']}"])
        ws.append([f"Summary: {s['summary']}"])
        ws.append([])
        ws.append(s["header"])
        photo_col = "ABCDEFGH"[len(s["header"]) - 1]
        for r in s["rows"]:
            ws.append(r)
            label = r[1] if s["header"][0] == "Bag" else r[0]
            ws.add_image(XLImage(_photo(label, n)), f"{photo_col}{ws.max_row}")
            n += 1
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if path:
        with open(path, "wb") as f:
            f.write(data)
    return data


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "sample_inventory.xlsx")
